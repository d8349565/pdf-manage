"""Write page content to one shared XY grid, retaining spatial rows and columns."""
from io import BytesIO
from math import ceil

from openpyxl.drawing.image import Image
from PIL import Image as RasterImage, ImageChops
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.pdf2excel.excel_builder import _write_cells
from app.pdf2excel.detector import EPSILON
from app.output import check_cancel


def _table_edges(table, vertical=False):
    start,end=(table.bbox[1],table.bbox[3]) if vertical else (table.bbox[0],table.bbox[2])
    count=table.row_count if vertical else table.col_count
    widths=table.row_heights if vertical else table.col_widths
    edges=[start]
    for width in widths or [(end-start)/count]*count:
        edges.append(edges[-1]+width)
    return edges


def _edges(values):
    groups=[]
    for value in sorted(values):
        if groups and value-groups[-1][0]<=EPSILON:
            groups[-1].append(value)
        else:
            groups.append([value])
    return [sum(group)/len(group) for group in groups]


def _grid(page, vertical=False):
    values=[] if vertical else [0.0,page['width']]
    for item in page['items']:
        if item['kind']=='table':
            values.extend(_table_edges(item['table'],vertical))
        else:
            values.extend(item['bbox'][1::2] if vertical else item['bbox'][::2])
    return _edges(values)


def _nearest(edges,value):
    return min(range(len(edges)),key=lambda index:abs(edges[index]-value))


def _backgrounds(page, cancel=None):
    """Transfer evidenced table tint to cells; residual paper must not cover text."""
    owners=[]
    for item in page['items']:
        check_cancel(cancel)
        if item['kind']=='text' and item.get('fill_colour'):
            owners.append((item['bbox'],item['fill_colour']))
        elif item['kind']=='table':
            owners.extend((cell.bbox,cell.style.fill_color) for cell in item['table'].cells if cell.style.fill_color)
    output=[]; fills=[]
    for item in page['items']:
        check_cancel(cancel)
        if item['kind']!='image' or item.get('image_role')!='background':
            output.append(item); continue
        image=RasterImage.open(BytesIO(item['data'])).convert('RGBA')
        width,height=image.size
        red,green,blue,_=image.split()
        left,top,right,bottom=item['bbox']; sx=(right-left)/width; sy=(bottom-top)/height
        for colour in {colour for _,colour in owners}:
            check_cancel(cancel)
            rgb=tuple(int(colour.lstrip('#')[offset:offset+2],16) for offset in (0,2,4))
            # Pillow applies these masks in native code. Keep byte-sized masks,
            # rather than a Python tuple and index list for every image pixel.
            mask=image.getchannel('A').point([0]+[255]*255)
            for channel,target in zip((red,green,blue),rgb):
                check_cancel(cancel)
                mask=ImageChops.multiply(mask,channel.point([255 if abs(value-target)<=16 else 0 for value in range(256)]))
            start=None; run_left=width; run_right=0
            for y in range(height+1):
                if y%32==0:
                    check_cancel(cancel)
                row=mask.crop((0,y,width,y+1)).tobytes() if y<height else b''
                if row.count(255)>=width*.3:
                    if start is None:start=y
                    run_left=min(run_left,row.find(b'\xff'))
                    run_right=max(run_right,row.rfind(b'\xff')+1)
                elif start is not None:
                    box=(left+run_left*sx,top+start*sy,left+run_right*sx,top+y*sy)
                    if any(value==colour and area[0]<box[2] and area[2]>box[0] and area[1]<box[3] and area[3]>box[1]
                           for area,value in owners):
                        fills.append(dict(kind='fill',bbox=box,colour=colour))
                        image.paste((0,0,0,0),(0,start,width,y),mask.crop((0,start,width,y)))
                    start=None; run_left=width; run_right=0
        white=red.point([255 if value>240 else 0 for value in range(256)])
        for channel in (green,blue):
            check_cancel(cancel)
            white=ImageChops.multiply(white,channel.point([255 if value>240 else 0 for value in range(256)]))
        image.paste((0,0,0,0),(0,0,width,height),white)
        check_cancel(cancel)
        stream=BytesIO(); image.save(stream,format='PNG')
        output.append(dict(item,data=stream.getvalue()))
    return dict(page,items=[*fills,*output])


def write_page(wb,page, *, cancel=None):
    """Use geometry for placement; item iteration never determines the next row."""
    page=_backgrounds(page,cancel)
    xs=_grid(page); ys=_grid(page,True)
    last_column=max(1,len(xs)-1)
    ws=wb.create_sheet(title=f"第{page['page_no']}页")
    ws.sheet_view.showGridLines=False
    for column,(left,right) in enumerate(zip(xs,xs[1:]),start=1):
        ws.column_dimensions[get_column_letter(column)].width=max(.1,(right-left)/5.625)
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=last_column)
    header=ws.cell(1,1,f"PDF 第 {page['page_no']} 页")
    header.font=Font(name='微软雅黑',size=10,bold=True,color='5A6571')
    header.fill=PatternFill(fill_type='solid',fgColor='EEF1F4')
    header.alignment=Alignment(vertical='center')
    ws.row_dimensions[1].height=24
    ws.freeze_panes='A2'

    # Native table boundaries and all other objects share these row intervals.
    # Split only Excel's oversized rows; compress empty gaps, never content.
    rows=[2]
    for top,bottom in zip(ys,ys[1:]):
        check_cancel(cancel)
        occupied=any(item['bbox'][1]<bottom and item['bbox'][3]>top for item in page['items'])
        height=max(1,bottom-top) if occupied else min(24,max(1,bottom-top))
        count=max(1,ceil(height/200))
        for row in range(rows[-1],rows[-1]+count):
            ws.row_dimensions[row].height=height/count
        rows.append(rows[-1]+count)

    priority={'table':0,'fill':1,'text':2,'image':3}
    for item in sorted(page['items'],key=lambda item:(priority[item['kind']],item['bbox'][1],item['bbox'][0])):
        check_cancel(cancel)
        left,top,right,bottom=item['bbox']
        column=_nearest(xs,left)+1; end_column=max(column,_nearest(xs,right))
        start_y=_nearest(ys,top); end_y=max(start_y+1,_nearest(ys,bottom))
        end_y=min(len(rows)-1,end_y)
        row=rows[start_y]; end_row=max(row,rows[end_y]-1)
        if item['kind']=='table':
            table=item['table']
            columns=[_nearest(xs,edge)+1 for edge in _table_edges(table)]
            table_rows=[rows[_nearest(ys,edge)] for edge in _table_edges(table,True)]
            _write_cells(ws,table,columns=columns,rows=table_rows,cancel=cancel)
        elif item['kind']=='fill':
            fill=PatternFill(fill_type='solid',fgColor=item['colour'].lstrip('#'))
            for cell_row in ws.iter_rows(min_row=row,max_row=end_row,min_col=column,max_col=end_column):
                check_cancel(cancel)
                for cell in cell_row:cell.fill=fill
        elif item['kind']=='text':
            # Glyph bounds are not a text container: allow the whitespace up
            # to the next field, while retaining the field's X/Y anchor.
            neighbours=[other['bbox'][0] for other in page['items'] if other is not item and other['kind'] in ('text','table')
                        and other['bbox'][0]>left+EPSILON and min(bottom,other['bbox'][3])-max(top,other['bbox'][1])>=min(bottom-top,other['bbox'][3]-other['bbox'][1])*.5]
            available=min(neighbours,default=page['width'])
            bands=[other['bbox'][2] for other in page['items'] if other['kind']=='fill'
                   and other['bbox'][0]<=left and other['bbox'][1]<=(top+bottom)/2<=other['bbox'][3]]
            if bands:
                available=min(available,min(bands))
            end_column=max(column,_nearest(xs,available))
            # OCR glyph boxes can overlap the following baseline. Containers
            # stop at the next field or table, so merges never consume its anchor.
            following=[other['bbox'][1] for other in page['items'] if other is not item and other['kind'] in ('text','table')
                       and other['bbox'][1]>top+EPSILON and min(available,other['bbox'][2])>max(left,other['bbox'][0])]
            if following:
                end_y=max(start_y+1,_nearest(ys,min(bottom,min(following))))
                end_row=max(row,rows[min(len(rows)-1,end_y)]-1)
            cell=ws.cell(row,column)
            text=item['text'].replace('\t','    ')
            # Retain both texts if a PDF genuinely overlays text in one grid
            # slot; do not overwrite earlier content while resolving overlaps.
            cell.value=(str(cell.value)+'\n'+text) if cell.value is not None else text
            cell.data_type='s'
            cell.font=Font(name='微软雅黑',size=max(6,min(72,item.get('size',11))),bold=item.get('bold',False),
                           color=item.get('colour','000000').lstrip('#'))
            if item.get('fill_colour'):
                cell.fill=PatternFill(fill_type='solid',fgColor=item['fill_colour'].lstrip('#'))
            multiline='\n' in str(cell.value)
            cell.alignment=Alignment(wrap_text=multiline,shrink_to_fit=not multiline,vertical='center')
            if end_column>column or end_row>row:
                ws.merge_cells(start_row=row,start_column=column,end_row=end_row,end_column=end_column)
        elif item['kind']=='image':
            image=Image(BytesIO(item['data']))
            image.width=(right-left)*4/3; image.height=(bottom-top)*4/3
            image.anchor=OneCellAnchor(
                _from=AnchorMarker(col=column-1,row=row-1,
                                   colOff=round(max(0,left-xs[column-1])*12700),
                                   rowOff=round(max(0,top-ys[start_y])*12700)),
                ext=XDRPositiveSize2D(cx=round((right-left)*12700),cy=round((bottom-top)*12700)))
            ws.add_image(image)

    ws.print_title_rows='1:1'
    ws.print_area=f'A1:{get_column_letter(last_column)}{max(1,rows[-1]-1)}'
    ws.page_setup.orientation='landscape' if page['width']>page['height'] else 'portrait'
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.oddFooter.center.text=f"原 PDF 第 {page['page_no']} 页"
