from __future__ import annotations

import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.pdf2excel.models import TableModel, CellModel, CellStyle

def build_workbook(tables: list[TableModel], *, pages: list[dict] | None = None) -> openpyxl.Workbook:
    """
    将 TableModel 列表转换为 openpyxl Workbook。
    默认每个表格一张工作表；提供 pages 时保留各页全部内容。
    """
    wb = openpyxl.Workbook()
    # 移除默认的 Sheet
    if wb.sheetnames:
        del wb[wb.sheetnames[0]]

    if pages is not None:
        from app.pdf2excel.page_builder import write_page
        for page in pages:
            write_page(wb, page)
        if not pages:
            wb.create_sheet(title="Sheet1")
        return wb
        
    if not tables:
        wb.create_sheet(title="Sheet1")
        return wb

    page_table_counts: dict[int, int] = {}
    for table in tables:
        p_no = table.page_no
        page_table_counts[p_no] = page_table_counts.get(p_no, 0) + 1
        sub_idx = page_table_counts[p_no]
        sheet_name = _sheet_name(sub_idx, table, len(tables))
        _write_table(wb, table, sheet_name)
        
    return wb

def _write_table(wb: openpyxl.Workbook, table: TableModel, sheet_name: str) -> None:
    """
    将单个表格写入新工作表。
    步骤：
    1. 创建工作表
    2. 设置列宽：pt_to_excel_width()
    3. 设置行高（直接映射磅值）
    4. 写入单元格值和数字格式
    5. 应用 merge_cells (当 col_span > 1 或 row_span > 1)
    6. 应用样式 (字体, 填充, 对齐, 边框)
    """
    ws = wb.create_sheet(title=sheet_name)
    
    # 2. 设置列宽
    if table.col_widths:
        for col_idx, width_pt in enumerate(table.col_widths, start=1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = pt_to_excel_width(width_pt)
            
    # 3. 设置行高
    if table.row_heights:
        for row_idx, height_pt in enumerate(table.row_heights, start=1):
            ws.row_dimensions[row_idx].height = height_pt

    _write_cells(ws, table)


def _write_cells(ws, table: TableModel, *, row_offset: int = 0,
                 columns: list[int] | None = None) -> None:
    """Write native table cells to their own sheet or a page's shared column grid."""
    # 记录已被合并覆盖的单元格，防止重复处理
    merged_cells = set()
    
    # 支持扁平 CellModel 列表或二维列表
    cell_list: list[CellModel] = []
    if table.cells:
        if isinstance(table.cells[0], (list, tuple)):
            for r_i, r_list in enumerate(table.cells):
                for c_i, c_obj in enumerate(r_list):
                    if c_obj is not None:
                        cell_list.append(c_obj)
        else:
            cell_list = table.cells

    for cell_model in cell_list:
        row_idx = cell_model.row + 1 + row_offset
        col_idx = columns[cell_model.col] if columns else cell_model.col + 1
        
        if (row_idx, col_idx) in merged_cells:
            continue
            
        cell = ws.cell(row=row_idx, column=col_idx)
        
        # 4. 写入值和数字格式
        val = cell_model.value if cell_model.value is not None else cell_model.text
        if val != "":
            cell.value = val
            if isinstance(val, str):
                # PDF 原文不是 Excel 公式（也包括 #N/A 等错误字面量）。
                cell.data_type = "s"
            
        if cell_model.number_format:
            cell.number_format = cell_model.number_format
            
        # 6. 应用样式
        if cell_model.style:
            _apply_style(cell, cell_model.style)
            
        # 5. 处理合并单元格
        end_col = (columns[cell_model.col + cell_model.col_span] - 1 if columns
                   else col_idx + cell_model.col_span - 1)
        if cell_model.row_span > 1 or end_col > col_idx:
            end_row = row_idx + cell_model.row_span - 1
            ws.merge_cells(
                start_row=row_idx,
                start_column=col_idx,
                end_row=end_row,
                end_column=end_col
            )
            
            # 将合并区域内的其余单元格标记为已处理
            for r in range(row_idx, end_row + 1):
                for c in range(col_idx, end_col + 1):
                    if (r, c) != (row_idx, col_idx):
                        merged_cells.add((r, c))

def pt_to_excel_width(pt: float) -> float:
    """
    将 PDF 磅值转换为 Excel 列宽单位。
    Excel 宽度 = 默认字体 (Calibri 11pt) 下 '0' 字符的个数。
    经验值: 96 DPI 下 1 Excel 单位 ≈ 5.625 pt。
    """
    return max(1.0, pt / 5.625)

def _clean_sheet_name(name: str) -> str:
    """清理 Excel 工作表名称（最大 31 个字符，不允许特殊字符）"""
    return re.sub(r'[\[\]:*?/\\]', '_', name)[:31]

def _sheet_name(index: int, table: TableModel, total: int) -> str:
    """
    生成唯一的工作表名称。
    如果 total == 1: 返回 'Sheet1'
    否则: 返回 'Page{page_no}_Table{n}' (兼容 Excel 最大长度 31)
    """
    if total == 1:
        return "Sheet1"
    
    page_no = getattr(table, 'page_no', None) or getattr(table, 'page_number', None) or 1
    name = f"Page{page_no}_Table{index}"
    return _clean_sheet_name(name)

def _color_hex_to_excel(hex_color: str | None) -> str | None:
    """转换颜色：从 '#FF0000' 到 openpyxl 的 'FF0000'"""
    if not hex_color:
        return None
    return hex_color.lstrip('#')

def _apply_style(cell, style: CellStyle) -> None:
    """
    将 CellStyle 应用于 openpyxl 单元格。
    映射:
    - font_name/size/bold/italic/color → Font()
    - fill_color → PatternFill(start_color=..., fill_type='solid')
    - horizontal/vertical → Alignment()
    - border_* → Border(Side(...))
    """
    if not style:
        return
        
    # 字体
    font_kwargs = {}
    if style.font_name:
        font_kwargs['name'] = style.font_name
    if style.font_size:
        font_kwargs['size'] = style.font_size
    if style.bold is not None:
        font_kwargs['bold'] = style.bold
    if style.italic is not None:
        font_kwargs['italic'] = style.italic
    if style.font_color:
        color = _color_hex_to_excel(style.font_color)
        if color:
            font_kwargs['color'] = color
            
    if font_kwargs:
        cell.font = Font(**font_kwargs)
        
    # 填充
    if style.fill_color:
        color = _color_hex_to_excel(style.fill_color)
        if color:
            cell.fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
            
    # 对齐方式
    align_kwargs = {}
    if style.horizontal_alignment:
        align_kwargs['horizontal'] = style.horizontal_alignment
    if style.vertical_alignment:
        align_kwargs['vertical'] = style.vertical_alignment
    if align_kwargs:
        cell.alignment = Alignment(wrap_text=True, **align_kwargs)
        
    # 边框
    border_kwargs = {}
    if style.border_left:
        border_kwargs['left'] = _side(style.border_left)
    if style.border_right:
        border_kwargs['right'] = _side(style.border_right)
    if style.border_top:
        border_kwargs['top'] = _side(style.border_top)
    if style.border_bottom:
        border_kwargs['bottom'] = _side(style.border_bottom)
        
    if border_kwargs:
        cell.border = Border(**border_kwargs)

def _side(style: str | None, color: str | None = None) -> Side:
    """将边框样式字符串转换为 openpyxl Side。"""
    if not style:
        return Side(style=None)
        
    color_val = _color_hex_to_excel(color)
    side_kwargs = {'style': style}
    if color_val:
        side_kwargs['color'] = color_val
        
    return Side(**side_kwargs)
