"""Orientation and spatial row contracts, independent of document wording."""
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
from PIL import Image
import pymupdf as p
import pytest

from app.ocr.word_export import W, convert_ocr_pdf_to_docx
from app.pdf2excel.converter import convert_pdf_to_excel
from app.pdf2excel.excel_builder import build_workbook
from app.documentlayout import prepare_export_page, _remove_page_rotation


def oriented_table(path, rotation=0, *, metadata=True, scanned=False):
    values=[['ID','DESCRIPTION','COUNT','RATE'],
            ['AX-01','Alpha', '12','7.25'],
            ['BZ-02','Beta', '18','9.50'],
            ['CY-03','Gamma', '6','0.125'],
            ['DZ-04','Delta', '14','1.75']]
    with p.open() as original:
        page=original.new_page(width=720,height=460)
        xs=[40,140,400,480,670]; ys=[100,136,172,208,244,280]
        page.insert_text((40,55),'Spatial export reference',fontsize=18)
        for x in xs:
            page.draw_line((x,ys[0]),(x,ys[-1]),width=.7)
        for y in ys:
            page.draw_line((xs[0],y),(xs[-1],y),width=.7)
        for row,entries in enumerate(values):
            for col,text in enumerate(entries):
                page.insert_text((xs[col]+6,ys[row]+23),text,fontsize=11)
        page.insert_text((40,325),'Footer must remain below the grid.',fontsize=11)
        if scanned:
            raw=page.get_text('dict')
            pixels=page.get_pixmap(dpi=120).tobytes('png')
            with p.open() as scan:
                target=scan.new_page(width=720,height=460)
                target.insert_image(target.rect,stream=pixels)
                for block in raw['blocks']:
                    for line in block.get('lines',()):
                        for span in line['spans']:
                            target.insert_text(span['origin'],span['text'],fontsize=span['size'],render_mode=3)
                carrier=scan.tobytes()
        else:
            carrier=original.tobytes()
    with p.open(stream=carrier,filetype='pdf') as source, p.open() as result:
        width,height=(460,720) if rotation in (90,270) else (720,460)
        page=result.new_page(width=width,height=height)
        page.show_pdf_page(page.rect,source,0,rotate=rotation)
        page.set_rotation(rotation if metadata else 0)
        result.save(path)
    return values


@pytest.mark.parametrize('rotation',[0,90,180,270])
@pytest.mark.parametrize('scanned',[False,True])
def test_rotated_all_content_keeps_horizontal_table_and_values(tmp_path,rotation,scanned):
    source=tmp_path/'source.pdf'; output=tmp_path/'output.xlsx'
    expected=oriented_table(source,rotation,scanned=scanned)
    original=source.read_bytes()
    result,_=convert_pdf_to_excel(str(source),str(output),include_all_content=True)
    assert source.read_bytes()==original
    assert len(result.tables)==1
    table=result.tables[0]
    assert (table.row_count,table.col_count)==(5,4)
    assert [[next(cell.text for cell in table.cells if cell.row==r and cell.col==c)
             for c in range(4)] for r in range(5)]==expected
    workbook=openpyxl.load_workbook(output)
    try:
        ws=workbook.active
        cells={cell.value:cell for row in ws for cell in row if cell.value is not None}
        assert cells['AX-01'].row==cells['Alpha'].row==cells[12].row==cells[7.25].row
        assert cells['AX-01'].column<cells['Alpha'].column<cells[12].column<cells[7.25].column
        assert ws.page_setup.orientation=='landscape'
    finally:
        workbook.close()


@pytest.mark.parametrize('rotation',[0,90,180,270])
def test_word_rotated_table_has_same_rows_and_landscape_paper(tmp_path,rotation):
    source=tmp_path/'source.pdf'; output=tmp_path/'output.docx'
    expected=oriented_table(source,rotation,scanned=True)
    before=source.read_bytes()
    state,_,_=convert_ocr_pdf_to_docx(source,output,False,None,lambda:False)
    assert state=='success' and source.read_bytes()==before
    with ZipFile(output) as archive:
        root=ET.fromstring(archive.read('word/document.xml'))
    tables=list(root.iter(f'{{{W}}}tbl'))
    assert len(tables)==1
    assert len(tables[0].findall(f'{{{W}}}tr'))==5
    assert len(tables[0].find(f'{{{W}}}tblGrid'))==4
    assert [[''.join(node.text or '' for node in cell.iter(f'{{{W}}}t'))
             for cell in row.findall(f'{{{W}}}tc')] for row in tables[0].findall(f'{{{W}}}tr')]==expected
    assert all(int(node.get(f'{{{W}}}w'))>int(node.get(f'{{{W}}}h')) for node in root.iter(f'{{{W}}}pgSz'))


def test_same_visual_row_does_not_become_vertical_list():
    picture=BytesIO(); Image.new('RGB',(12,12),'red').save(picture,format='PNG')
    page=dict(page_no=1,width=600,height=400,items=[
        dict(kind='text',text='Left field',bbox=(20,40,110,58),size=11),
        dict(kind='image',bbox=(280,36,300,60),data=picture.getvalue()),
        dict(kind='text',text='Right field',bbox=(390,41,490,59),size=11),
        dict(kind='text',text='Next row',bbox=(20,85,110,103),size=11)])
    workbook=build_workbook([],pages=[page])
    try:
        ws=workbook.active
        cells={cell.value:cell for row in ws for cell in row if cell.value is not None}
        assert cells['Left field'].row==cells['Right field'].row
        assert cells['Left field'].column<cells['Right field'].column
        assert cells['Next row'].row>cells['Left field'].row
    finally:
        workbook.close()


@pytest.mark.parametrize('rotation',[90,180,270])
def test_rotated_pixels_without_rotate_metadata_use_text_reading_axis(tmp_path,rotation):
    source=tmp_path/'pixel-rotation.pdf'; output=tmp_path/'pixel-rotation.xlsx'
    oriented_table(source,rotation,metadata=False,scanned=True)
    result,_=convert_pdf_to_excel(str(source),str(output),include_all_content=True)
    assert (result.tables[0].row_count,result.tables[0].col_count)==(5,4)
    assert next(cell.text for cell in result.tables[0].cells if cell.row==1 and cell.col==0)=='AX-01'


@pytest.mark.parametrize('rotation',[0,90,180,270])
@pytest.mark.parametrize('shifted_media',[False,True])
def test_baking_display_rotation_preserves_crop_pixels(rotation,shifted_media):
    with p.open() as doc:
        page=doc.new_page(width=200,height=300)
        if shifted_media:
            page.set_mediabox(p.Rect(10,20,210,320))
        page.draw_rect((25,35,175,265),fill=(.2,.6,.8),color=None)
        page.insert_text((45,100),'Crop reference',fontsize=13)
        media=page.mediabox
        page.set_cropbox(p.Rect(media.x0+30,40,media.x0+180,260))
        page.set_rotation(rotation)
        before=page.get_pixmap(dpi=72)
        _remove_page_rotation(page)
        after=page.get_pixmap(dpi=72)
        assert page.rotation==0
        assert (before.width,before.height)==(after.width,after.height)
        assert before.samples==after.samples


def test_rotated_header_does_not_rotate_horizontal_body_page():
    with p.open() as doc:
        page=doc.new_page(width=600,height=400)
        page.insert_text((50,180),'Vertical heading',fontsize=11,rotate=90)
        for row in range(5):
            page.insert_text((100,100+row*24),f'Horizontal body record {row} with stable reading direction',fontsize=11)
        before=page.get_pixmap(dpi=72).samples
        prepare_export_page(page)
        assert page.rect.width==600 and page.rect.height==400
        assert page.get_pixmap(dpi=72).samples==before


def shaded_sparse_table(path):
    values=[['ID','NAME','COUNT','AMOUNT'],['DEV-A','Alpha','12','0.011'],
            ['DEV-B','Beta','18',''],['DEV-C','Gamma','6','0.125']]
    xs=[40,140,400,480,670]; ys=[90,126,162,198,234]
    with p.open() as source:
        page=source.new_page(width=720,height=460)
        page.draw_rect((40,90,670,126),fill=(.15,.5,.6),color=None)
        for x in xs:
            page.draw_line((x,90),(x,234),color=(.6,.6,.6),width=.5,dashes='[3 3] 0')
        for y in ys:
            page.draw_line((40,y),(670,y),color=(.6,.6,.6),width=.5,dashes='[3 3] 0')
        for row,entries in enumerate(values):
            for col,text in enumerate(entries):
                if text:
                    page.insert_text((xs[col]+6,ys[row]+23),text,fontsize=11,
                                     color=(1,1,1) if row==0 else (0,0,0))
        raw=page.get_text('dict'); pixels=page.get_pixmap(dpi=144).tobytes('png')
    with p.open() as doc:
        page=doc.new_page(width=720,height=460); page.insert_image(page.rect,stream=pixels)
        for block in raw['blocks']:
            for line in block.get('lines',()):
                for span in line['spans']:
                    page.insert_text(span['origin'],span['text'],fontsize=span['size'],render_mode=3)
        doc.save(path)
    return values


def test_shaded_weak_header_keeps_data_columns_and_background(tmp_path):
    source=tmp_path/'shaded.pdf'; output=tmp_path/'shaded.xlsx'
    shaded_sparse_table(source)
    result,_=convert_pdf_to_excel(str(source),str(output),include_all_content=True)
    assert result.pages_exported==1
    workbook=openpyxl.load_workbook(output)
    try:
        ws=workbook.active
        cells={str(cell.value):cell for row in ws for cell in row if cell.value is not None}
        assert cells['ID'].column==cells['DEV-A'].column==cells['DEV-B'].column
        assert cells['NAME'].column==cells['Alpha'].column==cells['Beta'].column
        assert cells['DEV-A'].row==cells['Alpha'].row==cells['12'].row
        for label in ('ID','NAME','COUNT','AMOUNT'):
            assert cells[label].fill.fill_type=='solid'
            assert cells[label].font.color.rgb.endswith('FFFFFF')
    finally:
        workbook.close()
