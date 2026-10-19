import io

import openpyxl
import pymupdf
import pytest

from app.pdf2excel.analyzer import extract_drawings
from app.pdf2excel.converter import convert_pdf_to_excel
from app.pdf2excel.excel_builder import build_workbook
from app.pdf2excel.models import CellModel, TableModel
from app.pdf2excel.page_merger import merge_cross_page_tables
from app.pdf2excel.quality import assess_table
from app.pdf2excel.style_parser import build_cell_style
from app.pdf2excel.type_parser import parse_value
from app.output import OperationCancelled


def make_table(page, rows):
    return TableModel(page, (50, 50, 250, 150), len(rows), 2,
                      [CellModel(r, c, text=v, value=v)
                       for r, row in enumerate(rows) for c, v in enumerate(row)],
                      [50] * len(rows), [100, 100])


def roundtrip(wb):
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return openpyxl.load_workbook(stream)


def test_different_tables_stay_separate():
    tables = [make_table(1, [['Product', 'Price'], ['A', '10']]),
              make_table(2, [['Employee', 'Age'], ['B', '30']])]
    assert len(merge_cross_page_tables(tables)) == 2


def test_matching_first_row_never_drops_data():
    tables = [make_table(1, [['A', '10'], ['B', '20']]),
              make_table(2, [['A', '10'], ['C', '30']]),
              make_table(3, [['A', '10'], ['D', '40']])]
    merged = merge_cross_page_tables(tables)
    assert len(merged) == 1
    assert merged[0].row_count == 6
    assert [c.text for c in merged[0].cells if c.col == 0] == ['A', 'B', 'A', 'C', 'A', 'D']


@pytest.mark.parametrize('text', ['=1+1', '=SUM(A2:A3)', '#N/A', '#REF!'])
def test_literals_remain_text(text):
    cell = roundtrip(build_workbook([make_table(1, [[text, 'test']])])).active['A1']
    assert cell.value == text
    assert cell.data_type == 's'


@pytest.mark.parametrize('text', ['1234567890123456.78', '1,234,567,890,123,457',
                                 '¥1234567890123456.78', '1234567890123456.78%',
                                 '00123', '1234567890123456'])
def test_unsafe_numbers_roundtrip_as_original_text(text):
    value, fmt = parse_value(text)
    assert value == text
    table = make_table(1, [[text, '']])
    table.cells[0] = CellModel(0, 0, text=text, value=value, number_format=fmt)
    assert roundtrip(build_workbook([table])).active['A1'].value == text


@pytest.mark.parametrize(('text', 'fmt'), [('1,234.5678', '#,##0.0000'),
                                        ('12.3456%', '0.0000%'), ('123.4500', '0.0000')])
def test_decimal_places_preserved(text, fmt):
    assert parse_value(text)[1] == fmt


def test_pdf_rectangle_style_survives_excel_roundtrip():
    with pymupdf.open() as doc:
        page = doc.new_page()
        bbox = (50, 50, 150, 100)
        page.draw_rect(bbox, color=(0, 0, 0), fill=(1, 0, 0), width=1)
        lines, fills = extract_drawings(page)
        style = build_cell_style(bbox, None, [], lines, fills)
    table = make_table(1, [['red', '']])
    table.cells[0] = CellModel(0, 0, text='red', style=style)
    cell = roundtrip(build_workbook([table])).active['A1']
    assert cell.fill.fgColor.rgb[-6:] == 'FF0000'
    assert all(getattr(cell.border, side).style == 'medium' for side in ('left', 'right', 'top', 'bottom'))


def make_pdf(path, with_table=True):
    with pymupdf.open() as doc:
        for _ in range(2):
            page = doc.new_page(width=400, height=400)
            if with_table:
                for x in (50, 150, 250):
                    page.draw_line((x, 50), (x, 150))
                for y in (50, 100, 150):
                    page.draw_line((50, y), (250, y))
                for r, row in enumerate([['Name', 'Price'], ['ABC', '123']]):
                    for c, value in enumerate(row):
                        page.insert_text((55 + c * 100, 75 + r * 50), value)
            else:
                page.insert_text((50, 50), 'No tables here')
        doc.save(path)


def test_pipeline_defaults_to_separate_tables(tmp_path):
    source, output = tmp_path / 'input.pdf', tmp_path / 'out.xlsx'
    make_pdf(source)
    result, _ = convert_pdf_to_excel(str(source), str(output))
    assert len(result.tables) == 2
    wb = openpyxl.load_workbook(output)
    assert len(wb.worksheets) == 2
    for ws in wb.worksheets:
        assert list(ws.values) == [('Name', 'Price'), ('ABC', 123)]
    result, _ = convert_pdf_to_excel(str(source), str(output), merge_tables=True)
    assert len(result.tables) == 1
    assert list(openpyxl.load_workbook(output).active.values) == [
        ('Name', 'Price'), ('ABC', 123), ('Name', 'Price'), ('ABC', 123)]


@pytest.mark.parametrize('existing', [False, True])
def test_no_tables_does_not_write(tmp_path, existing):
    source, output = tmp_path / 'input.pdf', tmp_path / 'out.xlsx'
    make_pdf(source, False)
    if existing:
        output.write_bytes(b'original')
    result, _ = convert_pdf_to_excel(str(source), str(output))
    assert not result.tables
    assert output.read_bytes() == b'original' if existing else not output.exists()


@pytest.mark.parametrize('stage', ['整理', '生成'])
def test_late_cancel_preserves_output(tmp_path, stage):
    source, output = tmp_path / 'input.pdf', tmp_path / 'out.xlsx'
    make_pdf(source)
    output.write_bytes(b'original')
    cancelled = False
    def progress(cur, total, message):
        nonlocal cancelled
        cancelled = stage in message
    with pytest.raises(InterruptedError):
        convert_pdf_to_excel(str(source), str(output), progress_callback=progress,
                             cancel_check=lambda: cancelled)
    assert output.read_bytes() == b'original'
    assert not list(tmp_path.glob('.pdf2excel-*'))


@pytest.mark.parametrize('action', ['fail', 'cancel'])
def test_save_failure_or_cancellation_preserves_output(tmp_path, monkeypatch, action):
    source, output = tmp_path / 'input.pdf', tmp_path / 'out.xlsx'
    make_pdf(source)
    output.write_bytes(b'original')
    cancelled = False
    from app.pdf2excel.converter import ExcelWriter
    original_save = ExcelWriter.write_data
    def save(writer):
        nonlocal cancelled
        original_save(writer)
        if action == 'fail':
            raise OSError('simulated disk failure')
        cancelled = True
    monkeypatch.setattr(ExcelWriter, 'write_data', save)
    with pytest.raises(OSError if action == 'fail' else OperationCancelled):
        convert_pdf_to_excel(str(source), str(output), cancel_check=lambda: cancelled)
    assert output.read_bytes() == b'original'
    assert not list(tmp_path.glob('.pdf2excel-*'))


def test_merged_cells_count_as_covered_structure():
    table = make_table(1, [['title', ''], ['A', '10']])
    table.cells = [CellModel(0, 0, text='title', col_span=2), *table.cells[2:]]
    quality = assess_table(table)
    assert quality.structure_score == 1
    assert not quality.warnings
