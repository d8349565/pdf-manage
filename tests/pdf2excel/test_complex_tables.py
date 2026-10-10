from datetime import datetime

import openpyxl
import pymupdf
import pytest

from app.pdf2excel.converter import convert_pdf_to_excel
from app.pdf2excel.detector import _filter_watermarks


def contract_pdf(path, rotation=0, watermarks=True):
    source = pymupdf.open()
    page = source.new_page(width=1600, height=700)
    xs = [30, 60, 140, 300] + [300 + 60 * i for i in range(1, 22)]
    ys = [60, 80, 100, 120] + [120 + 30 * i for i in range(1, 18)]
    expected = {}

    def cell(row, col, text, rowspan=1, colspan=1):
        box = pymupdf.Rect(xs[col], ys[row], xs[col + colspan], ys[row + rowspan])
        page.draw_rect(box, width=.6, color=(0, 0, 0))
        page.insert_text((box.x0 + 3, box.y0 + 12), text, fontsize=8,
                         fontname='helv' if text.isascii() else 'china-s')
        expected[(row + 1, col + 1)] = text

    for col, title in enumerate(['No', 'Part', '零件名称']):
        cell(0, col, title, rowspan=3)
    cell(0, 3, '分项费用', colspan=21)
    cell(1, 3, 'Tooling', colspan=9)
    cell(1, 12, 'Transport', colspan=12)
    for col in range(3, 24):
        cell(2, col, f'Field{col}')
    for row in range(3, 20):
        for col in range(24):
            value = (str(row - 2) if col == 0 else f'A11-{row:04}BA' if col == 1
                     else f'涂料{row}' if col == 2 else '2026-07-01' if col == 3
                     else '0.0000' if col % 2 == 0 else '100')
            cell(row, col, value)
    if watermarks:
        for anchor in [(200, 400), (700, 400), (1200, 400)]:
            page.insert_text(anchor, '2026-CONFIDENTIAL', fontsize=28, color=(.8, .8, .8),
                             morph=(pymupdf.Point(anchor), pymupdf.Matrix(45)))
    if rotation:
        with pymupdf.open() as doc:
            w, h = (700, 1600) if rotation in (90, 270) else (1600, 700)
            target = doc.new_page(width=w, height=h)
            target.show_pdf_page(target.rect, source, 0, rotate=rotation)
            target.set_rotation(rotation)
            doc.save(path)
    else:
        source.save(path)
    source.close()
    return expected


@pytest.mark.parametrize('rotation', [0, 90, 180, 270])
def test_wide_merged_watermarked_contract(tmp_path, rotation):
    src, out = tmp_path / 'contract.pdf', tmp_path / 'contract.xlsx'
    expected = contract_pdf(src, rotation)
    original = src.read_bytes()
    result, _ = convert_pdf_to_excel(str(src), str(out))
    assert src.read_bytes() == original
    assert len(result.tables) == 1
    table = result.tables[0]
    assert (table.row_count, table.col_count) == (20, 24)
    ws = openpyxl.load_workbook(out).active
    assert {str(r) for r in ws.merged_cells.ranges} == {
        'A1:A3', 'B1:B3', 'C1:C3', 'D1:X1', 'D2:L2', 'M2:X2'}
    for (row, col), text in expected.items():
        value = ws.cell(row, col).value
        if row >= 4:
            if col == 1 or col > 4:
                assert value == float(text), (row, col, value, text)
            elif col == 4:
                assert value == datetime(2026, 7, 1)
            else:
                assert value == text
        else:
            assert value == text
    assert any('水印' in warning for warning in result.warnings)
    assert not any('CONFIDENTIAL' in str(cell.value) for row in ws for cell in row)


def test_watermark_filter_preserves_gray_labels_and_vertical_headers():
    spans = [dict(text='body', size=10, color=0, dir=(1, 0)),
             dict(text='Gray label', size=10, color=0xBBBBBB, dir=(1, 0)),
             dict(text='Vertical', size=28, color=0xBBBBBB, dir=(0, -1)),
             dict(text='Unique note', size=28, color=0xBBBBBB, dir=(.707, -.707))]
    assert _filter_watermarks(spans) == (spans, 0)


def test_single_pdf_span_across_columns_is_split_by_character(tmp_path):
    source, output = tmp_path / 'span.pdf', tmp_path / 'span.xlsx'
    with pymupdf.open() as doc:
        page = doc.new_page(width=200, height=200)
        for x in (50, 90, 130):
            page.draw_line((x, 50), (x, 150))
        for y in (50, 100, 150):
            page.draw_line((50, y), (130, y))
        page.insert_text((55, 75), 'AAA    BBB', fontname='cour', fontsize=10)
        page.insert_text((55, 125), '111    222', fontname='cour', fontsize=10)
        assert any('AAA    BBB' == span['text']
                   for block in page.get_text('dict')['blocks'] if block['type'] == 0
                   for line in block['lines'] for span in line['spans'])
        doc.save(source)
    convert_pdf_to_excel(str(source), str(output))
    assert list(openpyxl.load_workbook(output).active.values) == [('AAA', 'BBB'), (111, 222)]
