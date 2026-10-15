"""Focused all-content export checks; run only when simulation tests are requested."""
from io import BytesIO
from threading import Event

import openpyxl
import pymupdf as p
import pytest
from PIL import Image, ImageChops

from app.output import OperationCancelled
from app.pdf2excel.converter import convert_pdf_to_excel
from app.pdf2excel.excel_builder import build_workbook
from app.pdf2excel.models import CellModel, CellStyle, TableModel


def test_page_content_keeps_text_tables_and_blank_pages():
    table = TableModel(
        page_no=4, bbox=(50, 70, 250, 130), row_count=2, col_count=2,
        cells=[CellModel(0, 0, text='Merged header', col_span=2,
                         style=CellStyle(bold=True, fill_color='#336699')),
               CellModel(1, 0, text='00123', value='00123'),
               CellModel(1, 1, text='12.30', value=12.3, number_format='0.00')],
        row_heights=[30, 30], col_widths=[100, 100],
    )
    pages = [
        dict(page_no=4, width=300, height=400, items=[
            dict(kind='text', text='Title before table', bbox=(50, 20, 250, 40), size=14),
            dict(kind='table', table=table, bbox=table.bbox),
            dict(kind='text', text='=Original note', bbox=(50, 150, 250, 170), size=11),
        ]),
        dict(page_no=7, width=300, height=400, items=[
            dict(kind='text', text='No table, still exported', bbox=(30, 20, 250, 40), size=11),
        ]),
        dict(page_no=8, width=300, height=400, items=[]),
    ]
    workbook = build_workbook([table], pages=pages)
    try:
        assert workbook.sheetnames == ['第4页', '第7页', '第8页']
        ws = workbook.worksheets[0]
        cells = {cell.value: cell for row in ws for cell in row if cell.value is not None}
        assert cells['Title before table'].row < cells['Merged header'].row < cells['=Original note'].row
        assert cells['Merged header'].fill.fgColor.rgb.endswith('336699')
        assert cells['00123'].data_type == 's'
        assert cells[12.3].number_format == '0.00'
        assert cells['=Original note'].data_type == 's'
        assert any(cells['Merged header'].coordinate in merged for merged in ws.merged_cells.ranges)
        assert workbook.worksheets[1]['B2'].value == 'No table, still exported'
        assert workbook.worksheets[2].max_row == 1
        assert ws.max_row < 20
    finally:
        workbook.close()


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_all_content_keeps_visible_text_boxes_and_original_annotations(tmp_path, rotation):
    from app.pdfedits import apply_edits
    from app.workspace import PageEdit
    source, target = tmp_path / "annotated.pdf", tmp_path / "annotated.xlsx"
    with p.open() as doc:
        page = doc.new_page(width=400, height=600)
        page.insert_text((40, 50), "Original body")
        page.set_rotation(rotation)
        area = tuple(p.Rect(40, 100, 280, 150) * page.derotation_matrix)
        apply_edits(page, (PageEdit("freetext", area, text="Inserted note\nValue 123.45",
            annotation_id="zhixu-text-export", color=(0, 0, 0),
            details=(("font", "Helvetica"), ("fontsize", 12), ("rotation", rotation))),))
        doc.save(source)
    original = source.read_bytes()
    convert_pdf_to_excel(str(source), str(target), include_all_content=True)
    workbook = openpyxl.load_workbook(target)
    try:
        text = " ".join(str(cell.value) for row in workbook.active for cell in row if cell.value is not None)
        assert all(value in text for value in ("Original body", "Inserted note", "Value 123.45"))
    finally:
        workbook.close()
    assert source.read_bytes() == original
    with p.open(source) as doc:
        assert list(doc[0].annots())


def test_large_scan_uses_bounded_bitmap_and_reports_downsampling(tmp_path, monkeypatch):
    from app.documentimages import MAX_EXPORT_IMAGE_PIXELS, MAX_EXPORT_IMAGE_SIDE
    source, target = tmp_path / "large.pdf", tmp_path / "large.xlsx"
    stream = BytesIO()
    Image.new("RGB", (8, 8), "white").save(stream, format="PNG")
    with p.open() as doc:
        page = doc.new_page(width=7200, height=7200)
        page.insert_image(page.rect, stream=stream.getvalue())
        doc.save(source)
    render = p.Page.get_pixmap
    sizes = []
    def record(page, *args, **kwargs):
        image = render(page, *args, **kwargs)
        sizes.append((image.width, image.height))
        return image
    monkeypatch.setattr(p.Page, "get_pixmap", record)
    result, _ = convert_pdf_to_excel(str(source), str(target), include_all_content=True)
    assert sizes and all(w * h <= MAX_EXPORT_IMAGE_PIXELS and max(w, h) <= MAX_EXPORT_IMAGE_SIDE for w, h in sizes)
    assert any("降低分辨率" in warning for warning in result.warnings)


def test_background_processing_observes_cancellation_between_native_masks(monkeypatch):
    from app.pdf2excel.page_builder import _backgrounds
    stream = BytesIO()
    Image.new("RGBA", (200, 300), (50, 80, 120, 255)).save(stream, format="PNG")
    page = dict(items=[dict(kind="text", bbox=(0, 0, 100, 30), fill_colour="#325078"),
                       dict(kind="image", bbox=(0, 0, 200, 300), data=stream.getvalue(), image_role="background")])
    stop = Event()
    multiply = ImageChops.multiply
    def cancel_after_mask(*args):
        image = multiply(*args)
        stop.set()
        return image
    monkeypatch.setattr(ImageChops, "multiply", cancel_after_mask)
    with pytest.raises(OperationCancelled):
        _backgrounds(page, stop.is_set)


def test_cancel_during_excel_zip_write_keeps_existing_target(tmp_path, monkeypatch):
    from app.pdf2excel import converter
    source, target = tmp_path / "source.pdf", tmp_path / "target.xlsx"
    with p.open() as doc:
        doc.new_page().insert_text((40, 80), "Keep this content")
        doc.save(source)
    target.write_bytes(b"keep existing export")
    stop = Event()
    write_data = converter.ExcelWriter.write_data
    def cancel_before_first_zip_entry(writer):
        stop.set()
        return write_data(writer)
    monkeypatch.setattr(converter.ExcelWriter, "write_data", cancel_before_first_zip_entry)
    with pytest.raises(OperationCancelled):
        convert_pdf_to_excel(str(source), str(target), include_all_content=True, cancel_check=stop.is_set)
    assert target.read_bytes() == b"keep existing export"
    assert not list(tmp_path.glob(".pdf2excel-*"))
