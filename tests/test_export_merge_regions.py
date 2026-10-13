"""Editable exports must retain complete regions without overlapping cell merges."""
from itertools import combinations
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
import pymupdf as p
import pytest

from app.ocr.word_export import W, convert_ocr_pdf_to_docx
from app.pdf2excel.converter import convert_pdf_to_excel
from app.pdf2excel.excel_builder import build_workbook
from app.pdf2excel.models import CellModel, TableModel


def _nonoverlapping_merges(ws):
    for first, second in combinations(ws.merged_cells.ranges, 2):
        assert (first.max_col < second.min_col or second.max_col < first.min_col
                or first.max_row < second.min_row or second.max_row < first.min_row)


@pytest.mark.parametrize('reverse', [False, True])
def test_overlapping_text_bounds_keep_both_values_without_readonly_cell(tmp_path, reverse):
    items = [dict(kind='text', text='First field', bbox=(20, 40, 150, 65), size=11),
             dict(kind='text', text='Second field', bbox=(20, 58, 150, 80), size=11)]
    if reverse:
        items.reverse()
    workbook = build_workbook([], pages=[dict(page_no=1, width=500, height=300, items=items)])
    target = tmp_path / 'overlap.xlsx'
    workbook.save(target)
    workbook.close()
    loaded = openpyxl.load_workbook(target)
    try:
        ws = loaded.active
        values = {cell.value: cell for row in ws for cell in row if cell.value}
        assert values['First field'].row < values['Second field'].row
        _nonoverlapping_merges(ws)
    finally:
        loaded.close()


def test_text_next_to_native_table_does_not_merge_across_it():
    table = TableModel(1, (120, 40, 320, 100), 2, 2,
                       [CellModel(row, col, f'CELL-{row}-{col}', bbox=(120 + col*100, 40 + row*30, 220 + col*100, 70 + row*30))
                        for row in range(2) for col in range(2)], [30, 30], [100, 100])
    items = [dict(kind='table', table=table, bbox=table.bbox),
             dict(kind='text', text='Left field', bbox=(20, 42, 100, 58), size=11),
             dict(kind='text', text='Right field', bbox=(360, 42, 450, 58), size=11)]
    workbook = build_workbook([], pages=[dict(page_no=1, width=500, height=300, items=items)])
    try:
        values = {cell.value for row in workbook.active for cell in row if cell.value}
        assert {'Left field', 'Right field', 'CELL-0-0', 'CELL-0-1', 'CELL-1-0', 'CELL-1-1'} <= values
        _nonoverlapping_merges(workbook.active)
    finally:
        workbook.close()


def test_separate_measured_tables_are_not_joined_by_aligned_fields(tmp_path):
    def draw_grids(page):
        for left in (40, 320):
            for x in (left, left+100, left+200):page.draw_line((x, 90), (x, 198), width=.7)
            for y in (90, 126, 162, 198):page.draw_line((left, y), (left+200, y), width=.7)
    with p.open() as doc:
        page = doc.new_page(width=600, height=320)
        draw_grids(page)
        for left, prefix in ((40, 'L'), (320, 'R')):
            page.draw_rect((left, 90, left+200, 126), fill=(.15, .5, .6), color=None)
            values = [[prefix+'-CODE', prefix+'-NUM'], [prefix+'-A', '12'], [prefix+'-B', '18']]
            for row, entries in enumerate(values):
                for col, value in enumerate(entries):
                    page.insert_text((left+6+col*100, 113+row*36), value, fontsize=11)
        raw = page.get_text('dict'); image = page.get_pixmap(dpi=144).tobytes('png')
    source = tmp_path / 'separate.pdf'
    with p.open() as doc:
        page = doc.new_page(width=600, height=320)
        page.insert_image(page.rect, stream=image); draw_grids(page)
        for block in raw['blocks']:
            for line in block.get('lines', ()):
                for span in line['spans']:
                    page.insert_text(span['origin'], span['text'], fontsize=span['size'], render_mode=3)
        doc.save(source)
    result, _ = convert_pdf_to_excel(str(source), str(tmp_path/'separate.xlsx'), include_all_content=True)
    assert len(result.tables) == 2
    assert [(table.row_count, table.col_count) for table in result.tables] == [(3, 2), (3, 2)]


def mixed_rule_table(path, *, dark=True, sparse=False, tinted=False, centered=False, data_rows=3):
    values = [['CODE', 'DESCRIPTION', 'GROUP', 'STATUS'],
              ['A-01', 'Alpha', 'North', 'Ready'],
              ['B-02', 'Beta', 'South', '' if sparse else 'Pending'],
              ['C-03', 'Gamma', 'West', 'Done']]
    values = values[:data_rows+1]
    xs = [40, 140, 380, 470, 660]
    ys = [90, 126, 162, 198, 234]
    ys = ys[:len(values)+1]
    colour = (.15, .5, .6) if dark else (.88, .94, .97)
    with p.open() as original:
        page = original.new_page(width=720, height=420)
        page.insert_text((40, 55), 'Mixed border reference', fontsize=18)
        page.draw_rect((xs[0], ys[0], xs[-1], ys[1]), fill=colour, color=None)
        if tinted:
            for row in (row for row in (1, 3) if row<len(values)):
                page.draw_rect((xs[0], ys[row], xs[-1], ys[row+1]), fill=(.96, .985, 1), color=None)
        for x in xs:
            page.draw_line((x, ys[0]), (x, ys[-1]), color=(.6, .6, .6), width=.5, dashes='[3 3] 0')
        for y in ys:
            page.draw_line((xs[0], y), (xs[-1], y), color=(.6, .6, .6), width=.5, dashes='[3 3] 0')
        for row, entries in enumerate(values):
            for col, text in enumerate(entries):
                if text:
                    x=(xs[col]+xs[col+1]-p.get_text_length(text,fontsize=11))/2 if centered and row==0 else xs[col]+6
                    page.insert_text((x, ys[row] + 23), text, fontsize=11,
                                     color=(1, 1, 1) if row == 0 and dark else (0, 0, 0))
        raw = page.get_text('dict')
        image = page.get_pixmap(dpi=144).tobytes('png')
    with p.open() as doc:
        page = doc.new_page(width=720, height=420)
        page.insert_image(page.rect, stream=image)
        for block in raw['blocks']:
            for line in block.get('lines', ()):
                for span in line['spans']:
                    page.insert_text(span['origin'], span['text'], fontsize=span['size'], render_mode=3)
        # A valid local grid must not prevent examining the rest of the table.
        for x in xs[2:]:
            page.draw_line((x, ys[1]), (x, ys[-1]), width=.7)
        for y in ys[1:]:
            page.draw_line((xs[2], y), (xs[-1], y), width=.7)
        doc.save(path)
    return values


@pytest.mark.parametrize('dark,sparse,tinted,centered,data_rows', [(True, False, False, False, 3), (False, False, False, False, 3), (True, True, False, False, 3), (True, False, True, False, 3), (True, False, False, True, 3), (True, False, False, True, 2)])
def test_local_grid_does_not_split_whole_shaded_table(tmp_path, dark, sparse, tinted, centered, data_rows):
    source = tmp_path / 'mixed.pdf'
    expected = mixed_rule_table(source, dark=dark, sparse=sparse, tinted=tinted, centered=centered, data_rows=data_rows)
    before = source.read_bytes()
    target = tmp_path / 'mixed.xlsx'
    result, _ = convert_pdf_to_excel(str(source), str(target), include_all_content=True)
    assert len(result.tables) == 1
    table = result.tables[0]
    assert (table.row_count, table.col_count) == (len(expected), 4)
    actual = [['' for _ in range(4)] for _ in expected]
    for cell in table.cells:
        actual[cell.row][cell.col] = cell.text
        assert all((cell.style.border_left, cell.style.border_right, cell.style.border_top, cell.style.border_bottom))
        if cell.row == 0:
            assert cell.style.fill_color
            if dark:assert cell.style.font_color == '#FFFFFF'
        if tinted and cell.row in (1, 3):
            assert cell.style.fill_color
    assert actual == expected
    workbook = openpyxl.load_workbook(target)
    try:
        ws = workbook.active
        values = {cell.value: cell for row in ws for cell in row if cell.value}
        assert values['CODE'].column == values['A-01'].column == values['B-02'].column
        assert values['DESCRIPTION'].column == values['Alpha'].column == values['Beta'].column
        for row in expected:
            filled = [values[value] for value in row if value]
            assert len({cell.row for cell in filled}) == 1
        _nonoverlapping_merges(ws)
    finally:
        workbook.close()
    word = tmp_path / 'mixed.docx'
    convert_ocr_pdf_to_docx(source, word, False, None, lambda: False)
    with ZipFile(word) as archive:
        root = ET.fromstring(archive.read('word/document.xml'))
    tables = list(root.iter(f'{{{W}}}tbl'))
    assert len(tables) == 1
    assert len(tables[0].findall(f'{{{W}}}tr')) == len(expected)
    assert len(tables[0].find(f'{{{W}}}tblGrid')) == 4
    if tinted:
        for index in (1, 3):
            row = tables[0].findall(f'{{{W}}}tr')[index]
            assert all(cell.find(f'{{{W}}}tcPr/{{{W}}}shd') is not None for cell in row.findall(f'{{{W}}}tc'))
    assert source.read_bytes() == before
